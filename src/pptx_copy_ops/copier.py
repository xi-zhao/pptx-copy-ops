from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, Literal, Optional, Tuple

from pptx import Presentation
from pptx.slide import Slide
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

SlideCopyMode = Literal["shape", "part"]


@dataclass(frozen=True)
class SlideSpec:
    source_path: Path | str
    slide_index: int

    def resolved_path(self) -> Path:
        return Path(self.source_path).expanduser().resolve()


class SlideCopier:
    """Copy slides across PPTX files using `shape` or `part` mode."""

    _PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
    _REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    _REL_ATTR_PREFIX = f"{{{_REL_NS}}}"
    _SLD_LAYOUT_ID_TAG = f"{{{_PML_NS}}}sldLayoutId"
    _SLIDE_LAYOUT_RELTYPE = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
    )
    _SLIDE_RELTYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
    _SLIDE_MASTER_RELTYPE = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster"
    )
    _TAGS_RELTYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/tags"
    _NOTES_SLIDE_RELTYPE = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
    )
    _TAGS_TAG = f"{{{_PML_NS}}}tags"
    _NOTES_ID_TAG = f"{{{_PML_NS}}}notesId"
    _CUST_DATA_LST_TAG = f"{{{_PML_NS}}}custDataLst"

    def __init__(
        self,
        target_template: Path | str,
        clear_existing: bool = True,
        shape_copy_layout_slide_index: int | None = 0,
    ):
        self._target_template = Path(target_template).expanduser().resolve()
        self.presentation = Presentation(str(self._target_template))
        self._source_cache: Dict[str, Presentation] = {}
        self._shape_copy_layout = self._resolve_shape_copy_layout(
            shape_copy_layout_slide_index
        )
        if clear_existing:
            self._delete_all_slides()

    def copy_slides(self, specs: list[SlideSpec], mode: SlideCopyMode = "part") -> None:
        if mode not in {"shape", "part"}:
            raise ValueError(f"Unsupported copy mode: {mode}")
        for spec in specs:
            self.copy_slide(spec, mode=mode)

    def copy_slide(self, spec: SlideSpec, mode: SlideCopyMode = "part") -> Slide:
        source_slide = self._get_source_slide(spec)
        if mode == "part":
            return self._copy_slide_part(source_slide)
        if mode == "shape":
            return self._copy_slide_shape(source_slide)
        raise ValueError(f"Unsupported copy mode: {mode}")

    def save(self, output_path: Path | str) -> Path:
        destination = Path(output_path).expanduser().resolve()
        destination.parent.mkdir(parents=True, exist_ok=True)
        self.presentation.save(str(destination))
        return destination

    def _delete_all_slides(self) -> None:
        for index in range(len(self.presentation.slides) - 1, -1, -1):
            r_id = self.presentation.slides._sldIdLst[index].rId
            self.presentation.part.drop_rel(r_id)
            del self.presentation.slides._sldIdLst[index]

    def _resolve_shape_copy_layout(self, slide_index: int | None) -> Any | None:
        if slide_index is None or not len(self.presentation.slides):
            return None
        bounded_index = max(0, min(slide_index, len(self.presentation.slides) - 1))
        return self.presentation.slides[bounded_index].slide_layout

    def _get_source_slide(self, spec: SlideSpec) -> Slide:
        source_path = str(spec.resolved_path())
        prs = self._source_cache.get(source_path)
        if prs is None:
            prs = Presentation(source_path)
            self._source_cache[source_path] = prs
        if spec.slide_index < 0 or spec.slide_index >= len(prs.slides):
            raise IndexError(
                f"slide_index {spec.slide_index} out of range for {source_path} (total={len(prs.slides)})"
            )
        return prs.slides[spec.slide_index]

    def _collect_used_relationship_ids(self, source_slide: Slide) -> set[str]:
        used: set[str] = set()
        for element in source_slide._element.iter():
            for attr_name, attr_value in element.attrib.items():
                if attr_name.startswith(self._REL_ATTR_PREFIX) and attr_value:
                    used.add(attr_value)
        return used

    def _build_partname_template(self, source_partname: str) -> str:
        match = re.search(r"(\d+)(\.[^./]+)$", source_partname)
        if match:
            return f"{source_partname[:match.start(1)]}%d{match.group(2)}"
        ext_match = re.search(r"(\.[^./]+)$", source_partname)
        if ext_match:
            return f"{source_partname[:ext_match.start(1)]}%d{ext_match.group(1)}"
        return f"{source_partname}%d"

    def _allocate_partname(self, template: str, allocated: set[str]) -> PackURI:
        for index in range(1, 100000):
            candidate = template % index
            if candidate in allocated:
                continue
            allocated.add(candidate)
            return PackURI(candidate)
        raise RuntimeError(f"Cannot allocate partname from template: {template}")

    def _import_part_into_package(
        self,
        source_part: Any,
        target_package: Any,
        imported_parts: Dict[int, Any],
        allocated_partnames: set[str],
        keep_metadata_rels: bool = True,
    ) -> Any:
        cache_key = id(source_part)
        cached = imported_parts.get(cache_key)
        if cached is not None:
            return cached

        if str(getattr(source_part, "content_type", "")).startswith("image/"):
            # Avoid get_or_add_image_part() while importing detached part graphs:
            # next_image_partname() only scans rooted parts and can allocate
            # duplicate `/ppt/media/imageN.*` names for not-yet-related parts.
            template = self._build_partname_template(str(source_part.partname))
            new_partname = self._allocate_partname(template, allocated_partnames)
            source_blob = source_part.blob
            source_cls = source_part.__class__
            try:
                image_part = source_cls.load(
                    new_partname,
                    source_part.content_type,
                    target_package,
                    source_blob,
                )
            except Exception:
                image_part = Part.load(
                    new_partname,
                    source_part.content_type,
                    target_package,
                    source_blob,
                )
            imported_parts[cache_key] = image_part
            return image_part

        template = self._build_partname_template(str(source_part.partname))
        new_partname = self._allocate_partname(template, allocated_partnames)

        source_blob = source_part.blob
        source_cls = source_part.__class__
        try:
            imported_part = source_cls.load(
                new_partname,
                source_part.content_type,
                target_package,
                source_blob,
            )
        except Exception:
            imported_part = Part.load(
                new_partname,
                source_part.content_type,
                target_package,
                source_blob,
            )

        imported_parts[cache_key] = imported_part

        rid_map: Dict[str, str] = {}
        for rel in source_part.rels.values():
            if rel.reltype == self._NOTES_SLIDE_RELTYPE:
                continue
            if not keep_metadata_rels and rel.reltype == self._TAGS_RELTYPE:
                continue

            if rel.is_external:
                new_rid = imported_part.rels._add_relationship(
                    rel.reltype,
                    rel.target_ref,
                    is_external=True,
                )
                rid_map[rel.rId] = new_rid
                continue

            imported_target = self._import_part_into_package(
                rel.target_part,
                target_package,
                imported_parts,
                allocated_partnames,
                keep_metadata_rels=keep_metadata_rels,
            )
            new_rid = imported_part.rels._add_relationship(
                rel.reltype,
                imported_target,
                is_external=False,
            )
            rid_map[rel.rId] = new_rid

        if hasattr(imported_part, "_element"):
            root = imported_part._element
            orphan_nodes: list[Any] = []
            for element in root.iter():
                for attr_name in list(element.attrib.keys()):
                    if not attr_name.startswith(self._REL_ATTR_PREFIX):
                        continue
                    rid = element.attrib.get(attr_name)
                    if not rid:
                        continue
                    if rid in rid_map:
                        element.set(attr_name, rid_map[rid])
                    else:
                        if element.tag in {self._TAGS_TAG, self._NOTES_ID_TAG}:
                            orphan_nodes.append(element)
                            continue
                        try:
                            del element.attrib[attr_name]
                        except Exception:
                            pass
            for node in orphan_nodes:
                parent = node.getparent()
                if parent is None:
                    continue
                parent.remove(node)
                if parent.tag == self._CUST_DATA_LST_TAG and len(parent) == 0:
                    grandparent = parent.getparent()
                    if grandparent is not None:
                        grandparent.remove(parent)

        return imported_part

    def _copy_slide_shape(self, source_slide: Slide) -> Slide:
        target_layout = self._shape_copy_layout or (
            self.presentation.slide_layouts[6]
            if len(self.presentation.slide_layouts) > 6
            else self.presentation.slide_layouts[0]
        )
        target_slide = self.presentation.slides.add_slide(target_layout)

        for shape in list(target_slide.shapes):
            element = shape.element
            element.getparent().remove(element)

        used_rids = self._collect_used_relationship_ids(source_slide)
        rid_map: Dict[str, str] = {}
        imported_parts: Dict[int, Any] = {}
        allocated_partnames = {str(part.partname) for part in target_slide.part.package.iter_parts()}

        for rel in source_slide.part.rels.values():
            if rel.rId not in used_rids:
                continue
            if rel.reltype == self._NOTES_SLIDE_RELTYPE:
                continue
            if rel.is_external:
                new_rid = target_slide.part.rels._add_relationship(
                    rel.reltype,
                    rel.target_ref,
                    is_external=True,
                )
            else:
                imported_target = self._import_part_into_package(
                    rel.target_part,
                    target_slide.part.package,
                    imported_parts,
                    allocated_partnames,
                    keep_metadata_rels=True,
                )
                new_rid = target_slide.part.rels._add_relationship(
                    rel.reltype,
                    imported_target,
                    is_external=False,
                )
            rid_map[rel.rId] = new_rid

        for source_shape in source_slide.shapes:
            cloned = copy.deepcopy(source_shape.element)
            for element in cloned.iter():
                for attr_name in list(element.attrib.keys()):
                    if not attr_name.startswith(self._REL_ATTR_PREFIX):
                        continue
                    rid = element.attrib.get(attr_name)
                    if not rid:
                        continue
                    if rid in rid_map:
                        element.set(attr_name, rid_map[rid])
                    else:
                        try:
                            del element.attrib[attr_name]
                        except Exception:
                            pass
            target_slide.shapes._spTree.insert_element_before(cloned, "p:extLst")

        return target_slide

    def _copy_slide_part(self, source_slide: Slide) -> Slide:
        imported_parts: Dict[int, Any] = {}
        allocated_partnames = {str(part.partname) for part in self.presentation.part.package.iter_parts()}

        imported_slide_part = self._import_part_into_package(
            source_slide.part,
            self.presentation.part.package,
            imported_parts,
            allocated_partnames,
            keep_metadata_rels=True,
        )
        self._ensure_slide_master_registered(self.presentation, imported_slide_part)
        slide_rid = self.presentation.part.relate_to(imported_slide_part, self._SLIDE_RELTYPE)
        self.presentation.slides._sldIdLst.add_sldId(slide_rid)
        return imported_slide_part.slide

    def _ensure_slide_master_registered(self, presentation: Presentation, slide_part: Any) -> None:
        try:
            master_part = slide_part.slide_layout.part.slide_master.part
        except Exception:
            return

        master_rid = presentation.part.relate_to(master_part, self._SLIDE_MASTER_RELTYPE)
        sld_master_id_lst = presentation.part._element.get_or_add_sldMasterIdLst()
        for entry in sld_master_id_lst.sldMasterId_lst:
            if entry.rId == master_rid:
                return

        self._normalize_slide_layout_ids(presentation)

        used_ids: set[int] = set()
        for entry in sld_master_id_lst.sldMasterId_lst:
            raw_id = entry.get("id")
            if not raw_id:
                continue
            try:
                used_ids.add(int(raw_id))
            except Exception:
                continue

        for rel in presentation.part.rels.values():
            if rel.is_external or rel.reltype != self._SLIDE_MASTER_RELTYPE:
                continue
            master_root = getattr(getattr(rel, "target_part", None), "_element", None)
            if master_root is None:
                continue
            for node in master_root.iter():
                if node.tag != self._SLD_LAYOUT_ID_TAG:
                    continue
                raw_id = node.get("id")
                if not raw_id:
                    continue
                try:
                    used_ids.add(int(raw_id))
                except Exception:
                    continue

        next_id = max(used_ids) + 1 if used_ids else 2147483649
        while next_id in used_ids:
            next_id += 1

        new_entry = sld_master_id_lst._add_sldMasterId(rId=master_rid)
        new_entry.set("id", str(next_id))

    def _normalize_slide_layout_ids(self, presentation: Presentation) -> None:
        layout_nodes: list[Tuple[Any, Optional[int]]] = []
        max_layout_id = 2147483648
        reserved_master_ids: set[int] = set()
        sld_master_id_lst = presentation.part._element.find(f"{{{self._PML_NS}}}sldMasterIdLst")
        if sld_master_id_lst is not None:
            for entry in sld_master_id_lst:
                raw_id = entry.get("id")
                if not raw_id:
                    continue
                try:
                    reserved_master_ids.add(int(raw_id))
                except Exception:
                    continue

        for rel in presentation.part.rels.values():
            if rel.is_external or rel.reltype != self._SLIDE_MASTER_RELTYPE:
                continue
            master_part = getattr(rel, "target_part", None)
            root = getattr(master_part, "_element", None)
            if root is None:
                continue
            for node in root.iter():
                if node.tag != self._SLD_LAYOUT_ID_TAG:
                    continue
                raw_id = node.get("id")
                parsed_id: Optional[int] = None
                if raw_id:
                    try:
                        parsed_id = int(raw_id)
                        max_layout_id = max(max_layout_id, parsed_id)
                    except Exception:
                        parsed_id = None
                layout_nodes.append((node, parsed_id))

        used_layout_ids: set[int] = set()
        for node, parsed_id in layout_nodes:
            if (
                parsed_id is not None
                and parsed_id not in used_layout_ids
                and parsed_id not in reserved_master_ids
            ):
                used_layout_ids.add(parsed_id)
                continue
            max_layout_id += 1
            while max_layout_id in used_layout_ids or max_layout_id in reserved_master_ids:
                max_layout_id += 1
            node.set("id", str(max_layout_id))
            used_layout_ids.add(max_layout_id)


SlideSource = SlideSpec | tuple[Path | str, int]


def _coerce_slide_spec(source: SlideSource) -> SlideSpec:
    if isinstance(source, SlideSpec):
        return source
    source_path, slide_index = source
    return SlideSpec(source_path=source_path, slide_index=slide_index)


def copy_pptx_slides(
    *,
    target_template: Path | str,
    sources: Iterable[SlideSource],
    output_pptx: Path | str,
    mode: SlideCopyMode = "part",
    clear_existing: bool = True,
) -> Path:
    """Copy slides into a PPTX and save the result.

    This is the standalone public API for callers that do not need to manage a
    `SlideCopier` instance directly.
    """

    copier = SlideCopier(target_template=target_template, clear_existing=clear_existing)
    copier.copy_slides([_coerce_slide_spec(source) for source in sources], mode=mode)
    return copier.save(output_pptx)
