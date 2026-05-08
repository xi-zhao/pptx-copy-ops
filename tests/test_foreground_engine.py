from __future__ import annotations

import posixpath
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

from pptx_copy_ops.foreground import CopyPolicy, ForegroundCopyPolicy, ForegroundCopyRequest
from pptx_copy_ops.foreground.audit import audit_pptx_integrity
from pptx_copy_ops.foreground.engine import copy_foreground_slide
from pptx_copy_ops.foreground.inventory import DML_NS, PML_NS, REL_NS, collect_layer_inventory, q
from pptx_copy_ops.foreground.promoter import _copy_content_type

PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def _slide_texts(path: Path) -> list[str]:
    prs = Presentation(str(path))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
    return texts


def _rels_path(part_name: str) -> str:
    return posixpath.join(
        posixpath.dirname(part_name),
        "_rels",
        f"{posixpath.basename(part_name)}.rels",
    )


def test_copy_content_type_matches_default_extensions_case_insensitively() -> None:
    source_entries = {
        "[Content_Types].xml": f"""
<Types xmlns="{CONTENT_TYPES_NS}">
  <Default Extension="gif" ContentType="image/gif"/>
</Types>
""".strip().encode(),
    }
    output_entries = {
        "[Content_Types].xml": f"""
<Types xmlns="{CONTENT_TYPES_NS}">
  <Default Extension="xml" ContentType="application/xml"/>
</Types>
""".strip().encode(),
    }

    _copy_content_type(
        source_entries=source_entries,
        output_entries=output_entries,
        source_part="ppt/media/image1.GIF",
        copied_part="ppt/media/image2.GIF",
    )

    root = etree.fromstring(output_entries["[Content_Types].xml"])
    defaults = {
        default.get("Extension"): default.get("ContentType")
        for default in root.findall(f"{{{CONTENT_TYPES_NS}}}Default")
    }
    assert defaults["gif"] == "image/gif"


def _inject_master_shapes(path: Path) -> None:
    inventory = collect_layer_inventory(path, 0)
    temp_path = path.with_suffix(".inject.tmp")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(
        temp_path,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == inventory.master_part:
                root = etree.fromstring(data)
                sp_tree = root.find(f".//{q(PML_NS, 'spTree')}")
                assert sp_tree is not None
                sp_tree.append(
                    etree.fromstring(
                        f"""
<p:sp xmlns:p="{PML_NS}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:nvSpPr>
    <p:cNvPr id="9001" name="MASTER_BACKGROUND_MARKER"/>
    <p:cNvSpPr/>
    <p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="0" y="0"/><a:ext cx="12192000" cy="6858000"/></a:xfrm>
    <a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
  </p:spPr>
</p:sp>
""".strip()
                    )
                )
                sp_tree.append(
                    etree.fromstring(
                        f"""
<p:sp xmlns:p="{PML_NS}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:nvSpPr>
    <p:cNvPr id="9002" name="MASTER_FOREGROUND_MARKER"/>
    <p:cNvSpPr/>
    <p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="400000" y="200000"/><a:ext cx="2200000" cy="350000"/></a:xfrm>
    <a:solidFill><a:srgbClr val="111111"/></a:solidFill>
  </p:spPr>
  <p:txBody>
    <a:bodyPr/>
    <a:lstStyle/>
    <a:p><a:r><a:t>MASTER_FOREGROUND_MARKER</a:t></a:r></a:p>
  </p:txBody>
</p:sp>
""".strip()
                    )
                )
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(info, data)
    temp_path.replace(path)


def _ensure_content_type_override(
    root: etree._Element,
    *,
    part_name: str,
    content_type: str,
) -> None:
    normalized = f"/{part_name.lstrip('/')}"
    for override in root.findall(f"{{{CONTENT_TYPES_NS}}}Override"):
        if override.get("PartName") == normalized:
            override.set("ContentType", content_type)
            return
    override = etree.SubElement(root, f"{{{CONTENT_TYPES_NS}}}Override")
    override.set("PartName", normalized)
    override.set("ContentType", content_type)


def _inject_master_chart_with_style_dependency(path: Path) -> None:
    inventory = collect_layer_inventory(path, 0)
    master_rels_part = _rels_path(inventory.master_part)
    chart_part = "ppt/charts/chart99.xml"
    chart_rels_part = "ppt/charts/_rels/chart99.xml.rels"
    chart_style_part = "ppt/charts/style99.xml"
    chart_rel_id = "rIdForegroundChart"
    temp_path = path.with_suffix(".chart.tmp")

    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(
        temp_path,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zout:
        existing = set(zin.namelist())
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "[Content_Types].xml":
                root = etree.fromstring(data)
                _ensure_content_type_override(
                    root,
                    part_name=chart_part,
                    content_type="application/vnd.openxmlformats-officedocument.drawingml.chart+xml",
                )
                _ensure_content_type_override(
                    root,
                    part_name=chart_style_part,
                    content_type="application/vnd.ms-office.chartstyle+xml",
                )
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif info.filename == inventory.master_part:
                root = etree.fromstring(data)
                sp_tree = root.find(f".//{q(PML_NS, 'spTree')}")
                assert sp_tree is not None
                sp_tree.append(
                    etree.fromstring(
                        f"""
<p:graphicFrame xmlns:p="{PML_NS}"
                xmlns:a="{DML_NS}"
                xmlns:c="{CHART_NS}"
                xmlns:r="{REL_NS}">
  <p:nvGraphicFramePr>
    <p:cNvPr id="9100" name="MASTER_CHART_FOREGROUND"/>
    <p:cNvGraphicFramePr/>
    <p:nvPr/>
  </p:nvGraphicFramePr>
  <p:xfrm>
    <a:off x="500000" y="1000000"/>
    <a:ext cx="3000000" cy="1800000"/>
  </p:xfrm>
  <a:graphic>
    <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">
      <c:chart r:id="{chart_rel_id}"/>
    </a:graphicData>
  </a:graphic>
</p:graphicFrame>
""".strip()
                    )
                )
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif info.filename == master_rels_part:
                root = etree.fromstring(data)
                rel = etree.SubElement(root, f"{{{PKG_REL_NS}}}Relationship")
                rel.set("Id", chart_rel_id)
                rel.set(
                    "Type",
                    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart",
                )
                rel.set("Target", "../charts/chart99.xml")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(info, data)

        if chart_part not in existing:
            zout.writestr(
                chart_part,
                f"""<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<c:chartSpace xmlns:c="{CHART_NS}" xmlns:r="{REL_NS}">
  <c:date1904 val="0"/>
  <c:roundedCorners val="0"/>
  <c:style r:id="rIdStyle"/>
  <c:chart><c:plotArea/></c:chart>
</c:chartSpace>""",
            )
        if chart_rels_part not in existing:
            zout.writestr(
                chart_rels_part,
                f"""<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<Relationships xmlns="{PKG_REL_NS}">
  <Relationship Id="rIdStyle" Type="http://schemas.microsoft.com/office/2011/relationships/chartStyle" Target="style99.xml"/>
</Relationships>""",
            )
        if chart_style_part not in existing:
            zout.writestr(
                chart_style_part,
                "<?xml version='1.0' encoding='UTF-8' standalone='yes'?><cs:chartStyle xmlns:cs='http://schemas.microsoft.com/office/drawing/2012/chartStyle'/>",
            )
    temp_path.replace(path)


def test_foreground_promote_keeps_master_foreground_and_slide_text(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    target = tmp_path / "target.pptx"
    output = tmp_path / "output.pptx"

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(500000, 700000, 3000000, 400000).text = "SLIDE_FOREGROUND_MARKER"
    prs.save(source)
    _inject_master_shapes(source)
    Presentation().save(target)

    result = copy_foreground_slide(
        ForegroundCopyRequest(
            source_pptx=source,
            slide_index=0,
            target_pptx=target,
            policy=ForegroundCopyPolicy(copy_policy=CopyPolicy.FOREGROUND_PROMOTE),
        ),
        output_pptx=output,
    )

    texts = _slide_texts(output)
    assert "SLIDE_FOREGROUND_MARKER" in texts
    assert "MASTER_FOREGROUND_MARKER" in texts
    assert "MASTER_BACKGROUND_MARKER" not in "\n".join(texts)
    assert result.trace.promoted_elements
    assert result.trace.audit_issues == []
    assert audit_pptx_integrity(output) == []


def test_exact_part_copy_keeps_source_master_registration(tmp_path: Path) -> None:
    source = tmp_path / "source_exact.pptx"
    target = tmp_path / "target_exact.pptx"
    output = tmp_path / "output_exact.pptx"

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(500000, 700000, 3000000, 400000).text = "EXACT_COPY_MARKER"
    prs.save(source)
    Presentation().save(target)

    result = copy_foreground_slide(
        ForegroundCopyRequest(
            source_pptx=source,
            slide_index=0,
            target_pptx=target,
            policy=ForegroundCopyPolicy(copy_policy=CopyPolicy.EXACT_PART_COPY),
        ),
        output_pptx=output,
    )

    assert "EXACT_COPY_MARKER" in _slide_texts(output)
    assert result.trace.audit_issues == []

    with zipfile.ZipFile(output) as archive:
        presentation = etree.fromstring(archive.read("ppt/presentation.xml"))
        master_entries = presentation.findall(f"./{q(PML_NS, 'sldMasterIdLst')}/{q(PML_NS, 'sldMasterId')}")
        slide_nodes = presentation.findall(f"./{q(PML_NS, 'sldIdLst')}/{q(PML_NS, 'sldId')}")
        assert master_entries
        assert slide_nodes[-1].get(q(REL_NS, "id"))

    assert audit_pptx_integrity(output) == []


def test_foreground_promote_copies_promoted_chart_dependency_closure(tmp_path: Path) -> None:
    source = tmp_path / "source_chart.pptx"
    target = tmp_path / "target_chart.pptx"
    output = tmp_path / "output_chart.pptx"

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(500000, 700000, 3000000, 400000).text = "SLIDE_WITH_MASTER_CHART"
    prs.save(source)
    _inject_master_chart_with_style_dependency(source)
    Presentation().save(target)

    result = copy_foreground_slide(
        ForegroundCopyRequest(
            source_pptx=source,
            slide_index=0,
            target_pptx=target,
            policy=ForegroundCopyPolicy(copy_policy=CopyPolicy.FOREGROUND_PROMOTE),
        ),
        output_pptx=output,
    )

    assert result.trace.audit_issues == []
    assert any("MASTER_CHART_FOREGROUND" in item for item in result.trace.promoted_elements)
    assert audit_pptx_integrity(output) == []

    with zipfile.ZipFile(output) as archive:
        names = set(archive.namelist())
        assert "ppt/charts/chart99.xml" in names
        assert "ppt/charts/_rels/chart99.xml.rels" in names
        assert "ppt/charts/style99.xml" in names

        content_types = etree.fromstring(archive.read("[Content_Types].xml"))
        overrides = {
            override.get("PartName"): override.get("ContentType")
            for override in content_types.findall(f"{{{CONTENT_TYPES_NS}}}Override")
        }
        assert overrides["/ppt/charts/chart99.xml"] == (
            "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
        )
        assert overrides["/ppt/charts/style99.xml"] == "application/vnd.ms-office.chartstyle+xml"

        slide_parts = sorted(name for name in names if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
        promoted_slide = slide_parts[-1]
        rels = etree.fromstring(archive.read(_rels_path(promoted_slide)))
        chart_targets = [
            rel.get("Target")
            for rel in rels
            if rel.get("Type") == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
        ]
        assert "../charts/chart99.xml" in chart_targets
