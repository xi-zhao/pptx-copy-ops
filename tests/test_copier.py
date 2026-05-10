from __future__ import annotations

import base64
import os
import posixpath
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from pptx_copy_ops import copy_pptx_slides  # noqa: E402
from pptx_copy_ops.copier import SlideCopier, SlideSpec  # noqa: E402

_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jM5cAAAAASUVORK5CYII="
)


def _build_target_template(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if slide.shapes.title is not None:
        slide.shapes.title.text = "Target Template"
    prs.save(path)


def _build_source_deck(path: Path, title: str, body: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = body

    image_path = path.with_suffix(".png")
    image_path.write_bytes(_ONE_PIXEL_PNG)
    slide.shapes.add_picture(str(image_path), Inches(9.0), Inches(0.2), Inches(0.35), Inches(0.35))
    prs.save(path)
    image_path.unlink(missing_ok=True)


def _slide_texts(slide) -> list[str]:
    texts: list[str] = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = (shape.text_frame.text or "").strip()
        if text:
            texts.append(text)
    return texts


def _assert_no_duplicate_zip_members(pptx_path: Path) -> None:
    with zipfile.ZipFile(pptx_path) as archive:
        names = [item.filename for item in archive.infolist()]
    duplicates = [name for name, count in Counter(names).items() if count > 1]
    assert not duplicates, f"duplicate zip members detected: {sorted(duplicates)}"


def _assert_no_dangling_slide_relationship_refs(pptx_path: Path) -> None:
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel_prefix = f"{{{rel_ns}}}"
    with zipfile.ZipFile(pptx_path) as archive:
        names = set(archive.namelist())
        slide_parts = sorted(
            name for name in names if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
        for slide_part in slide_parts:
            rel_part = posixpath.join(
                posixpath.dirname(slide_part),
                "_rels",
                f"{posixpath.basename(slide_part)}.rels",
            )
            rel_ids: set[str] = set()
            if rel_part in names:
                rel_root = ET.fromstring(archive.read(rel_part))
                for rel in rel_root:
                    rel_id = rel.attrib.get("Id")
                    if rel_id:
                        rel_ids.add(rel_id)
            slide_root = ET.fromstring(archive.read(slide_part))
            for element in slide_root.iter():
                for attr_name, attr_value in element.attrib.items():
                    if not attr_name.startswith(rel_prefix):
                        continue
                    assert attr_value in rel_ids, (
                        f"dangling relationship reference {attr_name}={attr_value} in {slide_part}, "
                        f"missing from {rel_part}"
                    )


def _assert_master_ids_disjoint_from_layout_ids(pptx_path: Path) -> None:
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    with zipfile.ZipFile(pptx_path) as archive:
        presentation_root = ET.fromstring(archive.read("ppt/presentation.xml"))
        master_ids = {
            int(node.attrib["id"])
            for node in presentation_root.findall(f"./{{{p_ns}}}sldMasterIdLst/{{{p_ns}}}sldMasterId")
            if node.attrib.get("id")
        }

        layout_ids: set[int] = set()
        master_parts = sorted(
            name
            for name in archive.namelist()
            if name.startswith("ppt/slideMasters/slideMaster") and name.endswith(".xml")
        )
        for master_part in master_parts:
            master_root = ET.fromstring(archive.read(master_part))
            for node in master_root.findall(f".//{{{p_ns}}}sldLayoutId"):
                raw_id = node.attrib.get("id")
                if raw_id:
                    layout_ids.add(int(raw_id))

        overlap = sorted(master_ids & layout_ids)
        assert not overlap, f"sldMasterId/@id conflicts with sldLayoutId/@id: {overlap}"


def _registered_master_count(pptx_path: Path) -> int:
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    with zipfile.ZipFile(pptx_path) as archive:
        presentation_root = ET.fromstring(archive.read("ppt/presentation.xml"))
    return len(
        presentation_root.findall(f"./{{{p_ns}}}sldMasterIdLst/{{{p_ns}}}sldMasterId")
    )


def _slide_layout_target(pptx_path: Path, slide_index: int = 0) -> str:
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(pptx_path) as archive:
        presentation_root = ET.fromstring(archive.read("ppt/presentation.xml"))
        presentation_rels = ET.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
        slide_ids = presentation_root.findall(f"./{{{p_ns}}}sldIdLst/{{{p_ns}}}sldId")
        selected_slide_id = slide_ids[slide_index]
        selected_slide_rid = selected_slide_id.attrib[f"{{{r_ns}}}id"]
        slide_target = next(
            rel.attrib["Target"]
            for rel in presentation_rels.findall(f"{{{rel_ns}}}Relationship")
            if rel.attrib.get("Id") == selected_slide_rid
        )
        slide_part = posixpath.normpath(posixpath.join("ppt", slide_target))
        slide_rels_part = posixpath.join(
            posixpath.dirname(slide_part),
            "_rels",
            f"{posixpath.basename(slide_part)}.rels",
        )
        slide_rels = ET.fromstring(archive.read(slide_rels_part))
        return next(
            rel.attrib["Target"]
            for rel in slide_rels.findall(f"{{{rel_ns}}}Relationship")
            if rel.attrib.get("Type", "").endswith("/slideLayout")
        )


def test_part_mode_copy_two_slides_and_id_invariants(tmp_path: Path) -> None:
    target = tmp_path / "target.pptx"
    src1 = tmp_path / "source1.pptx"
    src2 = tmp_path / "source2.pptx"
    output = tmp_path / "out_part.pptx"

    _build_target_template(target)
    _build_source_deck(src1, title="Source A", body="Body A")
    _build_source_deck(src2, title="Source B", body="Body B")

    copier = SlideCopier(target)
    copier.copy_slides(
        [
            SlideSpec(src1, 0),
            SlideSpec(src2, 0),
        ],
        mode="part",
    )
    copier.save(output)

    generated = Presentation(output)
    assert len(generated.slides) == 2
    assert "Source A" in _slide_texts(generated.slides[0])
    assert "Source B" in _slide_texts(generated.slides[1])

    _assert_no_dangling_slide_relationship_refs(output)
    _assert_master_ids_disjoint_from_layout_ids(output)


def test_public_api_copies_slides_without_engine_dependency(tmp_path: Path) -> None:
    target = tmp_path / "target.pptx"
    src1 = tmp_path / "source1.pptx"
    src2 = tmp_path / "source2.pptx"
    output = tmp_path / "out_public_api.pptx"

    _build_target_template(target)
    _build_source_deck(src1, title="Standalone A", body="Body A")
    _build_source_deck(src2, title="Standalone B", body="Body B")

    destination = copy_pptx_slides(
        target_template=target,
        sources=[(src1, 0), (src2, 0)],
        output_pptx=output,
        mode="part",
    )

    assert destination == output.resolve()
    generated = Presentation(output)
    assert len(generated.slides) == 2
    assert "Standalone A" in _slide_texts(generated.slides[0])
    assert "Standalone B" in _slide_texts(generated.slides[1])
    assert _registered_master_count(output) > _registered_master_count(target)
    _assert_no_dangling_slide_relationship_refs(output)
    _assert_master_ids_disjoint_from_layout_ids(output)


def test_shape_mode_copy_two_slides_text_matches(tmp_path: Path) -> None:
    target = tmp_path / "target.pptx"
    src1 = tmp_path / "source1.pptx"
    src2 = tmp_path / "source2.pptx"
    output = tmp_path / "out_shape.pptx"

    _build_target_template(target)
    _build_source_deck(src1, title="Shape Source A", body="Shape Body A")
    _build_source_deck(src2, title="Shape Source B", body="Shape Body B")

    copier = SlideCopier(target)
    copier.copy_slides(
        [
            SlideSpec(src1, 0),
            SlideSpec(src2, 0),
        ],
        mode="shape",
    )
    copier.save(output)

    generated = Presentation(output)
    assert len(generated.slides) == 2
    assert "Shape Source A" in _slide_texts(generated.slides[0])
    assert "Shape Source B" in _slide_texts(generated.slides[1])

    _assert_no_dangling_slide_relationship_refs(output)


def test_shape_mode_inherits_target_template_first_slide_layout(tmp_path: Path) -> None:
    target = tmp_path / "target.pptx"
    source = tmp_path / "source.pptx"
    output = tmp_path / "out_shape_layout.pptx"

    _build_target_template(target)
    _build_source_deck(source, title="Shape Source", body="Shape Body")

    source_layout_target = _slide_layout_target(target)
    copier = SlideCopier(target)
    copier.copy_slide(SlideSpec(source, 0), mode="shape")
    copier.save(output)

    assert _slide_layout_target(output) == source_layout_target
    assert "Shape Source" in _slide_texts(Presentation(str(output)).slides[0])


def test_shape_mode_can_use_second_template_slide_layout(tmp_path: Path) -> None:
    target = tmp_path / "target.pptx"
    source = tmp_path / "source.pptx"
    output = tmp_path / "out_shape_second_layout.pptx"

    _build_target_template(target)
    target_prs = Presentation(str(target))
    target_prs.slides.add_slide(target_prs.slide_layouts[1])
    target_prs.save(target)
    _build_source_deck(source, title="Shape Source", body="Shape Body")

    second_layout_target = _slide_layout_target(target, slide_index=1)
    copier = SlideCopier(target, shape_copy_layout_slide_index=1)
    copier.copy_slide(SlideSpec(source, 0), mode="shape")
    copier.save(output)

    assert _slide_layout_target(output) == second_layout_target
    assert "Shape Source" in _slide_texts(Presentation(str(output)).slides[0])


def test_part_mode_copy_slide_with_multiple_images_has_unique_media_members(tmp_path: Path) -> None:
    from PIL import Image

    target = tmp_path / "target.pptx"
    source = tmp_path / "source_multi_image.pptx"
    output = tmp_path / "out_multi_image_part.pptx"

    _build_target_template(target)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    if slide.shapes.title is not None:
        slide.shapes.title.text = "Source Multi"
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "contains two different images"

    image1 = source.with_name("img_red.png")
    image2 = source.with_name("img_green.png")
    Image.new("RGB", (4, 4), (255, 0, 0)).save(image1)
    Image.new("RGB", (4, 4), (0, 255, 0)).save(image2)
    slide.shapes.add_picture(str(image1), Inches(8.8), Inches(0.2), Inches(0.35), Inches(0.35))
    slide.shapes.add_picture(str(image2), Inches(9.2), Inches(0.2), Inches(0.35), Inches(0.35))
    prs.save(source)
    image1.unlink(missing_ok=True)
    image2.unlink(missing_ok=True)

    copier = SlideCopier(target)
    copier.copy_slides([SlideSpec(source, 0)], mode="part")
    copier.save(output)

    _assert_no_duplicate_zip_members(output)


def test_cli_smoke(tmp_path: Path) -> None:
    target = tmp_path / "target.pptx"
    src1 = tmp_path / "source1.pptx"
    src2 = tmp_path / "source2.pptx"
    output = tmp_path / "out_cli.pptx"

    _build_target_template(target)
    _build_source_deck(src1, title="CLI A", body="CLI Body A")
    _build_source_deck(src2, title="CLI B", body="CLI Body B")

    env = os.environ.copy()
    env["PYTHONPATH"] = str(SRC) + os.pathsep + env.get("PYTHONPATH", "")
    cmd = [
        sys.executable,
        "-m",
        "pptx_copy_ops.cli",
        "--target",
        str(target),
        "--source",
        f"{src1}:1",
        "--source",
        f"{src2}:1",
        "--mode",
        "part",
        "--output",
        str(output),
    ]
    result = subprocess.run(cmd, text=True, capture_output=True, env=env)
    assert result.returncode == 0, f"stdout={result.stdout}\nstderr={result.stderr}"
    assert output.exists()
