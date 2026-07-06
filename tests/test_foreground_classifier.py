from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from pptx_copy_ops.foreground import (
    CopyPolicy,
    ElementClassification,
    ForegroundCopyPolicy,
    LayerName,
)
from pptx_copy_ops.foreground.classifier import classify_element
from pptx_copy_ops.foreground.inventory import collect_layer_inventory
from pptx_copy_ops.foreground.models import ElementRef


def test_foreground_contracts_are_importable() -> None:
    policy = ForegroundCopyPolicy(copy_policy=CopyPolicy.FOREGROUND_PROMOTE)

    assert policy.copy_policy == CopyPolicy.FOREGROUND_PROMOTE
    assert LayerName.SLIDE.value == "slide"
    assert ElementClassification.FOREGROUND.value == "foreground"


def test_collect_layer_inventory_finds_slide_layout_and_master(tmp_path: Path) -> None:
    pptx_path = tmp_path / "inventory.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Inventory Title"
    prs.save(pptx_path)

    inventory = collect_layer_inventory(pptx_path, 0)

    assert inventory.slide_part.startswith("ppt/slides/")
    assert inventory.layout_part.startswith("ppt/slideLayouts/")
    assert inventory.master_part.startswith("ppt/slideMasters/")
    assert any(
        item.layer == LayerName.SLIDE and item.text == "Inventory Title"
        for item in inventory.elements
    )


def test_collect_layer_inventory_rejects_invalid_slide_index(tmp_path: Path) -> None:
    pptx_path = tmp_path / "inventory.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(pptx_path)

    with pytest.raises(IndexError):
        collect_layer_inventory(pptx_path, 1)


def test_classifier_keeps_slide_local_content() -> None:
    element = ElementRef(
        layer=LayerName.SLIDE,
        part_name="ppt/slides/slide1.xml",
        element_index=2,
        tag="sp",
        text="客户价值",
        extents=(500000, 500000, 3000000, 600000),
    )

    decision = classify_element(element, ForegroundCopyPolicy())

    assert decision.classification == ElementClassification.FOREGROUND
    assert "slide-local" in decision.reason


def test_classifier_removes_slide_local_full_canvas_background() -> None:
    element = ElementRef(
        layer=LayerName.SLIDE,
        part_name="ppt/slides/slide1.xml",
        element_index=2,
        tag="pic",
        name="Background Image",
        text="",
        extents=(0, 0, 11520488, 6480174),
    )

    decision = classify_element(element, ForegroundCopyPolicy())

    assert decision.classification == ElementClassification.BACKGROUND
    assert "slide-local background" in decision.reason


def test_classifier_keeps_slide_local_background_for_exact_part_copy() -> None:
    element = ElementRef(
        layer=LayerName.SLIDE,
        part_name="ppt/slides/slide1.xml",
        element_index=2,
        tag="pic",
        name="Background Image",
        text="",
        extents=(0, 0, 11520488, 6480174),
    )

    decision = classify_element(
        element,
        ForegroundCopyPolicy(copy_policy=CopyPolicy.EXACT_PART_COPY),
    )

    assert decision.classification == ElementClassification.FOREGROUND


def test_classifier_keeps_small_master_logo_picture() -> None:
    element = ElementRef(
        layer=LayerName.MASTER,
        part_name="ppt/slideMasters/slideMaster1.xml",
        element_index=3,
        tag="pic",
        name="Logo Picture",
        extents=(200000, 200000, 1600000, 400000),
    )

    decision = classify_element(element, ForegroundCopyPolicy())

    assert decision.classification == ElementClassification.FOREGROUND
    assert "picture" in decision.reason


def test_classifier_removes_large_no_text_background_shape() -> None:
    element = ElementRef(
        layer=LayerName.MASTER,
        part_name="ppt/slideMasters/slideMaster1.xml",
        element_index=4,
        tag="sp",
        name="Background Rectangle",
        text="",
        extents=(0, 0, 12192000, 6858000),
    )

    decision = classify_element(element, ForegroundCopyPolicy())

    assert decision.classification == ElementClassification.BACKGROUND


def test_classifier_removes_master_placeholder_text() -> None:
    element = ElementRef(
        layer=LayerName.LAYOUT,
        part_name="ppt/slideLayouts/slideLayout1.xml",
        element_index=2,
        tag="sp",
        name="标题占位符 1",
        text="单击此处编辑母版标题样式",
        extents=(0, 0, 1000000, 300000),
    )

    decision = classify_element(element, ForegroundCopyPolicy())

    assert decision.classification == ElementClassification.PLACEHOLDER


def test_classifier_keeps_footer_when_policy_requires_it() -> None:
    element = ElementRef(
        layer=LayerName.MASTER,
        part_name="ppt/slideMasters/slideMaster1.xml",
        element_index=5,
        tag="sp",
        name="Footer Placeholder",
        text="Confidential",
        extents=(500000, 6500000, 4000000, 200000),
    )

    decision = classify_element(
        element,
        ForegroundCopyPolicy(
            copy_policy=CopyPolicy.EXACT_PART_COPY,
            keep_date_footer=True,
        ),
    )

    assert decision.classification == ElementClassification.FOREGROUND
